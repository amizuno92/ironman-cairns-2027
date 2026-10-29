from plan import SESSIONS, WEEKS
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter as L

F = "Arial"
H = Font(name=F, bold=True, color="FFFFFF", size=10)
HFILL = PatternFill("solid", fgColor="1F3A5F")
TITLE = Font(name=F, bold=True, size=16, color="1F3A5F")
SUB = Font(name=F, bold=True, size=11, color="1F3A5F")
N = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
BLUE = Font(name=F, size=10, color="0000FF", bold=True)
YEL = PatternFill("solid", fgColor="FFF2CC")
WRAP = Alignment(wrap_text=True, vertical="top")
thin = Side(style="thin", color="D9D9D9")
BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
SPORT_FILL = {"Swim": "DDEBF7", "Bike": "E2EFDA", "Run": "FCE4D6", "Strength": "EDE7F6", "Race": "FFD966", "Rest": "F2F2F2"}
KIND_FILL = {"recovery": "E2EFDA", "race": "FFD966", "taper": "DDEBF7", "reduced": "F2F2F2"}

wb = Workbook()


def header(ws, row, cols, widths=None):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font, cell.fill, cell.alignment, cell.border = H, HFILL, Alignment(wrap_text=True, vertical="center"), BOX
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[L(i)].width = w


# ================= START HERE =================
ws = wb.active
ws.title = "Start Here"
ws.column_dimensions["A"].width = 26
ws.column_dimensions["B"].width = 110
r = 1
ws["A1"] = "IRONMAN Cairns 2027 — 35-week plan (Andrew)"; ws["A1"].font = TITLE
ws["A2"] = "12 Oct 2026 → 13 Jun 2027 · 80/20 polarized · built 11 Oct 2026"; ws["A2"].font = Font(name=F, italic=True, size=10, color="595959")
r = 4


def block(title, rows):
    global r
    ws.cell(row=r, column=1, value=title).font = SUB
    r += 1
    for k, v in rows:
        a = ws.cell(row=r, column=1, value=k); a.font = BOLD; a.alignment = WRAP
        b = ws.cell(row=r, column=2, value=v); b.font = N; b.alignment = WRAP
        r += 1
    r += 1


block("Key dates", [
    ("Wed–Fri 21–23 Oct 2026", "Wedding week (Fri 23 Oct). Train Mon–Thu, Fri–Sat off, easy ride Sun 25 Oct."),
    ("Mon 26 Oct 2026", "Running returns: walk-run stage 1 (you passed 30 pain-free hops)."),
    ("Week 14 (11–17 Jan)", "Checkpoint: tibia + London go/no-go; FTP & CSS retests."),
    ("Sun 21 Mar 2027", "Gold Coast T100, 100km distance (2km / 80km / 18km). Your first triathlon and the half-distance tune-up."),
    ("Sun 25 Apr 2027", "London Marathon at Z2 (HR ≤149). Fly Wed 21 Apr; no road bike until Mon 10 May."),
    ("Sun 13 Jun 2027", "IRONMAN Cairns 🏁"),
])
block("Goal: sub-11", [
    ("Target", "~10:55: swim 1:05 · T1/T2 10' · bike 5:30 · run 4:10. Realistic today: ~11:30–12:15."),
    ("What changes", "Volume stays the same (it's already near your 20h cap, and the tibia limits running). From week 15, IM power on long rides "
     "and Thursday blocks rises from 66–71% to 70–75% FTP. Every bike target scales with FTP, so retests raise them automatically."),
    ("Biggest levers", "FTP to ~265–275W, a fitted TT bike by ~January with long rides in aero, and 90–100 g carbs/h on the bike."),
    ("Decision", "Final race goal is set after the T100 test (21 Mar). Track FTP checkpoints on the Sub-11 tab."),
])
block("Recommendations (your Q3 & Q10)", [
    ("Lead-up races", "Yes, race once. Do the Gold Coast T100 100km race on Sun 21 Mar. It's at home, 12 weeks before Cairns, and the "
     "only chance to rehearse open-water mass starts, transitions and race fuelling. Race it controlled (see Calendar wk 23), not "
     "flat out, because London is 5 weeks later. If the wk-14 tibia check isn't clean, enter the Olympic distance on Sat 20 Mar instead. "
     "Optional: a local sprint in February for transition practice. Skip a separate 70.3: Geelong is on the same weekend as the T100."),
    ("London block (21 Apr–9 May)", "Last ride Tue 20 Apr; the last long ride is Sat 17 Apr. In London, swim 3–4×/wk and lift 2×/wk. "
     "Use the gym's spin bike or Wattbike by heart rate (wks 29–30): it keeps your cycling legs and costs nothing. Marathon recovery: "
     "no running for 5 days, first run only if pain-free hopping. You return with ~5 weeks to Cairns, which is enough for a 2-week "
     "IM-specific block, a last big weekend and a 2-week taper."),
])
block("How the plan works", [
    ("Periodisation", "Transition (wks 1–2) → Base 1–3 (3–13) → Build 1–3 (15–27) → London (28–30) → Peak (31–33) → Taper (34–35). "
     "There's a recovery or test week roughly every 4th week: wks 6, 10, 14, 18, 22, 24."),
    ("10% rule", "No 'load' week's total hours go up more than 10% on the previous load week. Run km follow the same cap once "
     "continuous running is established (wk 9). Recovery, race and taper weeks are lower by design, so the first load week after "
     "one is compared with the last load week before it. Walk-run weeks 3–8 follow a return-to-run protocol instead, because 10% of zero is nothing. "
     "Weekly Summary shows the % change for every load week."),
    ("80/20", "Most time is Z1–Z2. Hard work is Z4+ (bike VO2 and threshold, CSS swims, cruise intervals) with almost nothing in Z3. "
     "At this training age, time above Z2 sits at about 10% of training time; IM-power work sits at the top of Z2 power. "
     "Weekly Summary shows the high-intensity % for each week."),
    ("Week shape", "Mon swim (or easy spin in open-water weeks) + easy run · Tue bike intervals, then Strength A later (hard day) · Wed CSS swim + run quality · Thu endurance bike + run · "
     "Fri long swim + Strength B · Sat long ride + brick · Sun long run + easy spin or open-water swim. Max 3 swims a week. Swims at midday; bikes and solo runs in the morning; on bike days the run is in the evening; gym in the evening (see the Calendar's When column). Weekdays ≤2h, two sessions on most days."),
    ("Tibia rules", "Pain ≤2/10 during a run and nothing the next morning. If either fails, stop, take 2 days off and drop back one stage. "
     "Run on flat, soft surfaces at a slightly higher cadence, and use the soleus, tibialis and hop work in Strength A (bone loading). "
     "Run intensity only starts in wk 16, and only if the wk-14 checkpoint is clean. Any focal, pinpoint bone pain means seeing a sports doctor."),
    ("Bike", "Zwift until your helmet, computer and Varia arrive. Then move the long rides outdoors at the same power. "
     "Practise the TT position as soon as the TT bike arrives (fit first)."),
    ("Swim", "50m pool, so every rep is ≥50m and Apple Watch counts correctly. Open water starts in wk 17: buy the wetsuit by early February "
     "and do some non-wetsuit swims too."),
    ("Nutrition", "Gut training. Long rides go 60→100 g carbs/h over the season (you've already managed 90 g/h running). "
     "Long runs 60–90 g/h. Sodium 500–1000 mg/h depending on sweat. Rehearse your exact race products at the T100 and London."),
    ("Heat", "Cairns is warm and humid. Heat sessions start in wk 31: Zwift rides with no fan or an extra layer."),
    ("Retests", "CSS: wks 6, 14, 18. FTP ramp: wks 6, 14, 22. After a test, update the yellow cells on the Zones tab. "
     "Calendar watt targets assume FTP 242W; scale them by your new FTP ÷ 242."),
    ("Logging", "Tick the 'Done' column on the Calendar. Log FTP changes in Strava as 'FTP SET: [W]'."),
])

# ================= ZONES =================
z = wb.create_sheet("Zones")
for c, w in zip("ABCDEF", [30, 16, 16, 16, 30, 10]):
    z.column_dimensions[c].width = w
z["A1"] = "Training zones"; z["A1"].font = TITLE
z["A2"] = "Yellow cells are inputs. Update them after each test; every zone recalculates."; z["A2"].font = Font(name=F, italic=True, size=10)
inputs = [("Bike FTP (W)", 242), ("Bike LTHR (bpm)", 161), ("Bike max HR", 175), ("Run LTHR (bpm)", 167), ("Run max HR", 184),
          ("Swim LTHR (bpm)", 152), ("Swim max HR", 168), ("Swim CSS (sec per 100m)", 94)]
for i, (k, v) in enumerate(inputs, 4):
    z.cell(row=i, column=1, value=k).font = BOLD
    c = z.cell(row=i, column=2, value=v); c.font = BLUE; c.fill = YEL; c.border = BOX
z["C4"] = "Source: your ramp test (Sep 2026)"; z["C11"] = "Source: CSS test (Sep 2026). 94s = 1:34/100m"
z["C5"] = "Source: you (Oct 2026)"; z["C4"].font = z["C11"].font = z["C5"].font = Font(name=F, size=9, italic=True, color="595959")

row = 13


def ztable(title, cols, rows):
    global row
    z.cell(row=row, column=1, value=title).font = SUB
    row += 1
    for i, c in enumerate(cols, 1):
        cell = z.cell(row=row, column=i, value=c); cell.font, cell.fill = H, HFILL
    row += 1
    for rr in rows:
        for i, v in enumerate(rr, 1):
            cell = z.cell(row=row, column=i, value=v); cell.font = N; cell.border = BOX
        row += 1
    row += 1


ztable("Bike — power (% FTP, Coggan) & HR (% LTHR, Friel)", ["Zone", "Power low (W)", "Power high (W)", "HR range", "Use", "RPE"], [
    ["Z1 Recovery", 0, "=ROUND($B$4*0.55,0)", '="< "&ROUND($B$5*0.81,0)', "Recovery spins, warm-ups", "1–2"],
    ["Z2 Endurance", "=ROUND($B$4*0.56,0)", "=ROUND($B$4*0.75,0)", '=ROUND($B$5*0.81,0)&"–"&ROUND($B$5*0.89,0)', "Most riding. IM power = 70–75% from wk 15 (row below)", "3–4"],
    ["  IM race power (sub-11)", "=ROUND($B$4*0.70,0)", "=ROUND($B$4*0.75,0)", '="≤ "&ROUND($B$5*0.90,0)', "Ironman target; base weeks 1–14 use 66–71%", "4"],
    ["Z3 Tempo (avoid)", "=ROUND($B$4*0.76,0)", "=ROUND($B$4*0.90,0)", '=ROUND($B$5*0.90,0)&"–"&ROUND($B$5*0.93,0)', "Grey zone; T100 race effort only", "5–6"],
    ["Z4 Threshold", "=ROUND($B$4*0.91,0)", "=ROUND($B$4*1.05,0)", '=ROUND($B$5*0.94,0)&"–"&ROUND($B$5*0.99,0)', "Threshold intervals", "7–8"],
    ["Z5 VO2max", "=ROUND($B$4*1.06,0)", "=ROUND($B$4*1.20,0)", '="≥ "&ROUND($B$5*1.00,0)', "3–6' intervals", "8–9"],
])
ztable("Run — HR (% LTHR, Friel). Run by HR while the tibia rebuilds; pace is a guide only.", ["Zone", "HR low", "HR high", "% LTHR", "Use", "RPE"], [
    ["Z1 Recovery", "", "=ROUND($B$7*0.85,0)-1", "<85%", "Walk-run, shakeouts", "2–3"],
    ["Z2 Aerobic", "=ROUND($B$7*0.85,0)", "=ROUND($B$7*0.89,0)", "85–89%", "Easy and long runs. HARD CAP. London + IM run", "3–4"],
    ["Z3 Tempo (avoid)", "=ROUND($B$7*0.90,0)", "=ROUND($B$7*0.94,0)", "90–94%", "Grey zone", "5–6"],
    ["Z4 Threshold", "=ROUND($B$7*0.95,0)", "=ROUND($B$7*0.99,0)", "95–99%", "Cruise intervals (from wk 16)", "7"],
    ["Z5 VO2", "=$B$7", "=$B$8", "100%+", "Not used this season", "8–9"],
])
ztable("Swim — pace from CSS (HR from your wrist is unreliable in water; use pace + RPE)", ["Zone", "Pace fast (/100m)", "Pace slow (/100m)", "HR (guide)", "Use", "RPE"], [
    ["Z1 Recovery", '=INT(($B$11+18)/60)&":"&TEXT(MOD($B$11+18,60),"00")', "—", '="< "&ROUND($B$9*0.85,0)', "Easy / recovery", "2"],
    ["Z2 Aerobic / IM pace", '=INT(($B$11+8)/60)&":"&TEXT(MOD($B$11+8,60),"00")', '=INT(($B$11+14)/60)&":"&TEXT(MOD($B$11+14,60),"00")', '=ROUND($B$9*0.85,0)&"–"&ROUND($B$9*0.89,0)', "Aerobic sets; IM target pace = CSS+10s", "3–4"],
    ["Z3 Tempo (minimise)", '=INT(($B$11+3)/60)&":"&TEXT(MOD($B$11+3,60),"00")', '=INT(($B$11+7)/60)&":"&TEXT(MOD($B$11+7,60),"00")', '=ROUND($B$9*0.90,0)&"–"&ROUND($B$9*0.94,0)', "Rarely", "5–6"],
    ["Z4 CSS / threshold", '=INT(($B$11-2)/60)&":"&TEXT(MOD($B$11-2,60),"00")', '=INT(($B$11+2)/60)&":"&TEXT(MOD($B$11+2,60),"00")', '=ROUND($B$9*0.95,0)&"–"&ROUND($B$9*0.99,0)', "CSS sets", "7"],
    ["Z5 Speed", "—", '=INT(($B$11-4)/60)&":"&TEXT(MOD($B$11-4,60),"00")', '="≥ "&$B$9', "50s fast", "8–9"],
])

# ================= CALENDAR =================
cal = wb.create_sheet("Calendar")
cols = ["Week", "Date", "Day", "Phase", "Sport", "Session", "Duration (min)", "Distance (km)", "Main set / details",
        "Target zone (HR / Power / Pace)", "RPE", "Fuel", "High-intensity min (Z4+)", "Done", "Notes", "When"]
header(cal, 1, cols, [6, 11, 5, 11, 9, 26, 9, 9, 80, 38, 7, 30, 10, 7, 25, 14])
cal.freeze_panes = "G2"
for i, s in enumerate(SESSIONS, 2):
    vals = [s["wk"], s["date"], s["day"], s["phase"], s["sport"], s["title"], s["dur"], s["dist"], s["details"],
            s["target"], s["rpe"], s["fuel"], s["high"], "", s["note"], s.get("when", "")]
    fill = PatternFill("solid", fgColor=SPORT_FILL[s["sport"]])
    for j, v in enumerate(vals, 1):
        c = cal.cell(row=i, column=j, value=v)
        c.font = BOLD if s["sport"] == "Race" else N
        c.alignment = WRAP; c.border = BOX
        if j in (5, 6):
            c.fill = fill
    cal.cell(row=i, column=2).number_format = "ddd d mmm yy"
    cal.cell(row=i, column=8).number_format = "0.0"
last = len(SESSIONS) + 1
cal.auto_filter.ref = f"A1:P{last}"
dv = DataValidation(type="list", formula1='"✓,Partial,Missed"', allow_blank=True)
cal.add_data_validation(dv); dv.add(f"N2:N{last}")
cal.conditional_formatting.add(f"A2:P{last}", FormulaRule(formula=['$N2="✓"'], font=Font(name=F, color="7F7F7F")))

# ================= WEEKLY SUMMARY =================
wsum = wb.create_sheet("Weekly Summary", 1)
cols = ["Week", "Week starts", "Phase", "Week type", "Focus", "Total h", "Swim h", "Bike h", "Run h", "Strength h", "Race h",
        "Swim km", "Run km", "Δ hours vs prev load wk", "Δ run km vs prev load wk", "High-intensity % (Z4+)"]
header(wsum, 1, cols, [6, 11, 11, 10, 60, 8, 7, 7, 7, 9, 7, 8, 8, 12, 12, 12])
wsum.freeze_panes = "F2"
CR = f"Calendar!$A$2:$A${last}"
prev_load = prev_run = None
for i, n in enumerate(sorted(WEEKS), 2):
    w = WEEKS[n]
    vals = [n, w["start"], w["phase"], w["kind"], w["focus"]]
    for j, v in enumerate(vals, 1):
        c = wsum.cell(row=i, column=j, value=v); c.font = N; c.alignment = WRAP; c.border = BOX
    wsum.cell(row=i, column=2).number_format = "d mmm yy"
    sp = lambda s: f'SUMIFS(Calendar!$G$2:$G${last},{CR},$A{i},Calendar!$E$2:$E${last},"{s}")'
    f = {6: f"=SUMIFS(Calendar!$G$2:$G${last},{CR},$A{i})/60", 7: f"={sp('Swim')}/60", 8: f"={sp('Bike')}/60",
         9: f"={sp('Run')}/60", 10: f"={sp('Strength')}/60", 11: f"={sp('Race')}/60",
         12: f'=SUMIFS(Calendar!$H$2:$H${last},{CR},$A{i},Calendar!$E$2:$E${last},"Swim")',
         13: f'=SUMIFS(Calendar!$H$2:$H${last},{CR},$A{i},Calendar!$E$2:$E${last},"Run")',
         16: f"=IFERROR(SUMIFS(Calendar!$M$2:$M${last},{CR},$A{i})/(({sp('Swim')})+({sp('Bike')})+({sp('Run')})),0)"}
    if w["kind"] == "load" and prev_load:
        f[14] = f"=F{i}/F{prev_load}-1"
    if w["kind"] == "load" and n > 9 and prev_run:
        f[15] = f"=M{i}/M{prev_run}-1"
    for j in range(6, 17):
        c = wsum.cell(row=i, column=j, value=f.get(j, "—" if j in (14, 15) else None))
        c.font = N; c.border = BOX
        c.number_format = "0.0%" if j in (14, 15, 16) else "0.0"
        if j in (14, 15): c.alignment = Alignment(horizontal="right")
    if w["kind"] in KIND_FILL:
        for j in range(1, 6):
            wsum.cell(row=i, column=j).fill = PatternFill("solid", fgColor=KIND_FILL[w["kind"]])
    if w["kind"] == "load":
        prev_load = i
        if n >= 9: prev_run = i
lr = len(WEEKS) + 1
wsum.conditional_formatting.add(f"N2:O{lr}", FormulaRule(formula=["AND(ISNUMBER(N2),N2>0.1)"], fill=PatternFill("solid", fgColor="F4B6B6")))
note = lr + 2
wsum.cell(row=note, column=1, value="Δ columns only apply to 'load' weeks and compare with the previous load week; red would flag >10%. "
          "Run-km cap applies from week 9, once continuous running is established. Totals are live SUMIFS from the Calendar, so editing a "
          "session updates this table.").font = Font(name=F, size=9, italic=True)

ch = BarChart(); ch.type = "col"; ch.grouping = "stacked"; ch.overlap = 100
ch.title = "Weekly hours by sport"; ch.y_axis.title = "Hours"; ch.x_axis.title = "Week"
data = Reference(wsum, min_col=7, max_col=11, min_row=1, max_row=lr)
ch.add_data(data, titles_from_data=True); ch.set_categories(Reference(wsum, min_col=1, min_row=2, max_row=lr))
for s, col in zip(ch.series, ["5B9BD5", "70AD47", "ED7D31", "8064A2", "FFC000"]):
    s.graphicalProperties.solidFill = col; s.graphicalProperties.line.solidFill = col
ch.height, ch.width = 9, 26
wsum.add_chart(ch, f"A{note + 2}")

# ================= STRENGTH =================
st = wb.create_sheet("Strength")
header(st, 1, ["Phase (weeks)", "Session", "Exercises (sets × reps)", "Load / notes"], [22, 14, 90, 45])
from plan import STR
rows = [
    ("Base (1–14)", "A — Tue, after bike intervals", STR[("A", "base")], "Wks 1–5: 3×8–10 at RPE 7 to learn; wks 7–13: 4×5–6 heavy (2 reps in reserve). Hops: start double-leg pogos; single-leg from wk 8 only if pain-free."),
    ("Base (1–14)", "B — Thu", STR[("B", "base")], "Moderate–heavy, 2 RIR. Keep 48h before Sat long ride."),
    ("Build (15–30)", "A — Tue, after bike intervals", STR[("A", "build")], "Maintain max strength in fewer sets. Explosive intent on hops/jumps."),
    ("Build (15–30)", "B — Fri", STR[("B", "build")], "Deliberately light on legs: big ride on Saturday."),
    ("Peak (31–33)", "A — Tue, after bike intervals", STR[("A", "peak")], "20–30 min. Nothing to failure."),
    ("Peak (31–33)", "B — Fri", STR[("B", "peak")], "Mobility focus."),
    ("Taper (34)", "A — Tue only", "Short maintenance as Peak A at ~60% load.", "Last strength session: Tue 1 Jun. None in race week."),
    ("Recovery weeks", "A & B", "Same exercises, ~2/3 sets, 20–30% lighter.", "Leave feeling fresher than you started."),
]
for i, rr in enumerate(rows, 2):
    for j, v in enumerate(rr, 1):
        c = st.cell(row=i, column=j, value=v); c.font = N; c.alignment = WRAP; c.border = BOX

# ================= RACE DAY =================
rd = wb.create_sheet("Race Day")
rd.column_dimensions["A"].width = 24; rd.column_dimensions["B"].width = 100
rd["A1"] = "Race execution"; rd["A1"].font = TITLE
rr = 3
for title, items in [
    ("IRONMAN Cairns — Sun 13 Jun 2027", [
        ("Swim 3.8km", "Target 1:42–1:46/100m (~64–68'). Start wide and calm for the first 200m, find feet, sight every 6–9 strokes. Z2 effort; you should exit feeling you held back."),
        ("Bike 180km", "70–75% FTP NP: 170–182W at today's 242W, ~188–200W if FTP reaches 268W. Set the final number after the wk-22 test and the T100. HR ≤145. Cap rises at 220W. Sub-11 needs ~5:30; a fallback day is ~5:50–6:05. "
         "Eat on a timer: 90–100 g carbs/h, 700–1000 mg sodium/h, 600–900 ml/h; more fluid if hot."),
        ("Run 42.2km", "Km 0–10: HR 135–145, run slower than it feels you should (≈6:00–6:15/km). Km 10–30: HR ≤149, walk 10–20s at every aid station. "
         "Km 30–42: race whatever is left. Carbs 60–90 g/h (gels + cola late), sodium 500–800 mg/h."),
        ("Sub-11 target", "Swim ~1:05 + T1/T2 ~10' + bike ~5:30 + run ~4:10 ≈ 10:55. Needs FTP ~265–275W and a fitted TT bike. See the Sub-11 tab."),
        ("Fallback", "If the T100 test isn't met: bike ~5:50 at 68–70% FTP, run ~4:15 → ~11:15–11:30. A well-run 11:20 beats a sub-11 attempt that falls apart at 25km."),
        ("Carb load", "Fri–Sat: 8–10 g carbs/kg (≈640–800 g/day), low fibre. Race breakfast 3h before: 120–160 g carbs."),
    ]),
    ("Gold Coast T100 — Sun 21 Mar 2027 (2km / 80km / 18km)", [
        ("Purpose", "IM rehearsal: wetsuit, mass start, transitions, nutrition. Not an all-out effort, because London is 5 weeks later."),
        ("Targets", "Swim 1:40–1:44/100m · Bike 175–190W (72–78% FTP), HR ≤150 · Run HR ≤149 for 12km, then ≤155 if the tibia is quiet."),
    ]),
    ("London Marathon — Sun 25 Apr 2027", [
        ("Purpose", "A long Z2 training run in a great atmosphere. Not a PB."),
        ("Targets", "HR cap 149 (≈5:35–5:50/km → ~4:00). Fuel 60–90 g/h exactly as at Cairns. Stop if you get focal shin pain."),
        ("After", "5 days no running · swim, walk, gym bike · first run Sat 1 May only if pain-free hopping."),
    ]),
]:
    rd.cell(row=rr, column=1, value=title).font = SUB; rr += 1
    for k, v in items:
        a = rd.cell(row=rr, column=1, value=k); a.font = BOLD; a.alignment = WRAP
        b = rd.cell(row=rr, column=2, value=v); b.font = N; b.alignment = WRAP
        rr += 1
    rr += 1


# ================= SUB-11 =================
sx = wb.create_sheet("Sub-11")
for c, w in zip("ABCDEFG", [30, 14, 14, 14, 16, 44, 4]):
    sx.column_dimensions[c].width = w
sx["A1"] = "Sub-11 tracker"; sx["A1"].font = TITLE
sx["A2"] = "Enter your test results in the yellow cells. Status updates automatically."; sx["A2"].font = Font(name=F, italic=True, size=10)
header(sx, 4, ["FTP checkpoint", "Date", "Target FTP (W)", "Actual FTP (W)", "Status", "Notes"])
cps = [("Baseline (Sep ramp test)", "Sep 2026", 242, 242, "Source: your Zwift ramp test."),
       ("Week 6 ramp test", "Sat 21 Nov 2026", 252, None, "+10W: realistic after 6 weeks of base."),
       ("Week 14 ramp test", "Tue 12 Jan 2027", 260, None, "On track for sub-11 if at or above target."),
       ("Week 22 ramp test", "Tue 9 Mar 2027", 268, None, "Sets T100 and Cairns race power."),
       ("Race FTP needed", "13 Jun 2027", 268, None, "~270W lets 190–200W NP be ~72% FTP.")]
for i, (n, d, t, a, note) in enumerate(cps, 5):
    sx.cell(row=i, column=1, value=n).font = BOLD
    sx.cell(row=i, column=2, value=d).font = N
    c = sx.cell(row=i, column=3, value=t); c.font = N
    c = sx.cell(row=i, column=4, value=a); c.font = BLUE; c.fill = YEL; c.border = BOX
    sx.cell(row=i, column=5, value=f'=IF(D{i}="","—",IF(D{i}>=C{i},"On track",IF(D{i}>=C{i}-5,"Close","Behind")))').font = BOLD
    sx.cell(row=i, column=6, value=note).font = N
    for j in range(1, 7): sx.cell(row=i, column=j).border = BOX
sx.cell(row=11, column=1, value="Race power from latest FTP").font = SUB
sx.cell(row=12, column=1, value="Latest FTP (W)").font = BOLD
sx["B12"] = "=IFERROR(LOOKUP(2,1/(D5:D9<>\"\"),D5:D9),242)"; sx["B12"].font = N
sx.cell(row=13, column=1, value="IM NP at 70% / 75% (W)").font = BOLD
sx["B13"] = "=ROUND(B12*0.70,0)"; sx["C13"] = "=ROUND(B12*0.75,0)"
sx.cell(row=14, column=1, value="W/kg at 75% (80kg)").font = BOLD
sx["B14"] = "=ROUND(C13/80,2)"
for r in (13, 14):
    for col in "BC": sx[f"{col}{r}"].font = N
sx.cell(row=16, column=1, value="Split targets").font = SUB
header(sx, 17, ["Leg", "Sub-11", "Fallback", "", "", "Key requirement"])
splits = [("Swim 3.8km", "1:05", "1:07", "1:42/100m incl. sighting; open water from wk 17."),
          ("T1 + T2", "0:10", "0:12", "Practise at the T100; Cairns transitions are long."),
          ("Bike 180km", "5:30", "5:50", "190–200W NP in aero; TT bike fitted by ~January."),
          ("Run 42.2km", "4:10", "4:15", "HR ≤149 to 30km; walk aid stations; tibia healthy."),
          ("Total", "10:55", "11:24", "")]
for i, (l, a, f, k) in enumerate(splits, 18):
    for j, v in enumerate([l, a, f, "", "", k], 1):
        c = sx.cell(row=i, column=j, value=v); c.font = BOLD if l == "Total" else N; c.border = BOX
sx.cell(row=24, column=1, value="T100 go/no-go test (Sun 21 Mar)").font = SUB
for i, t in enumerate(["Bike 80km at ~200W NP (or ≥74% of FTP) with HR ≤150 in aero.",
                       "Run 18km at HR ≤149 for 12km, finishing strong with no fade and no shin pain.",
                       "Both met → race Cairns for sub-11. Either missed → race the fallback plan (~11:15–11:30)."], 25):
    c = sx.cell(row=i, column=1, value=t); c.font = N

wb._sheets = [wb[n] for n in ["Start Here", "Weekly Summary", "Calendar", "Zones", "Sub-11", "Strength", "Race Day"]]
wb.active = 0
wb.save("Ironman_Cairns_2027_Plan.xlsx")
print(wb.sheetnames, len(SESSIONS))
